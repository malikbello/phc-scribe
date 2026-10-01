"""Stratum U (ADR-0019): run the danger-sign rules on the real Uyo antenatal records.

    python scripts/check_uyo.py                         # download, verify, print the mapping
    python scripts/check_uyo.py --confirmed-by "Rahmat"  # also score against diagnoses

The mapping (bench/uyo.py) must be confirmed by the clinical adviser before any score
is computed, so the script refuses to score without --confirmed-by. The data file is
pinned by SHA-256 and kept under data/private/ (git-ignored); only aggregate counts
are printed.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import urllib.request
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "serve"), str(ROOT)]

from bench.uyo import SYMPTOM_FINDINGS, URINE_PROTEIN_LABELS, UyoRecord, evaluate  # noqa: E402

# OSF j9zh8, "Maternal, Neonatal and Child Health Datasets for Spatiotemporal Data
# Analytics" (CC-BY-4.0); file "MNCH Datasets_UyoSenDist.xlsx", downloaded 1 Oct 2026.
URL = "https://osf.io/download/mrjg9/"
SHA256 = "7f20ee4d9fa6c2443b524fa2a3843e8f6fa81cd185be57af85113f31b00eb81f"
PATH = ROOT / "data" / "private" / "uyo" / "mnch_uyo.xlsx"
OUT = ROOT / "results" / "uyo"


def download() -> Path:
    if not PATH.exists():
        PATH.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=120) as r:  # noqa: S310 (fixed https URL)
            PATH.write_bytes(r.read())
    got = hashlib.sha256(PATH.read_bytes()).hexdigest()
    if got != SHA256:
        sys.exit(f"checksum mismatch for {PATH.name}: expected {SHA256}, got {got}")
    return PATH


def load(path: Path) -> list[UyoRecord]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)["Martanal"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    col = {name: header.index(name) for name in header}
    return [
        UyoRecord(
            patient=str(r[col["Patient Number"]]),
            symptom=str(r[col["Symptom"]] or ""),
            diagnosis=str(r[col["Diagnose"]] or ""),
            bp=str(r[col["Blood Pressure (mm Hg)"]] or ""),
            temperature=str(r[col["Temperature (Degree C)"]] or ""),
        )
        for r in rows
        if any(v is not None for v in r)
    ]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--confirmed-by", help="clinical adviser who confirmed the mapping")
    args = parser.parse_args()
    records = load(download())
    labels = Counter(" ".join(r.symptom.split()).lower() for r in records)
    print(f"{len(records)} antenatal records; symptom labels and their mapping:")
    for label, n in labels.most_common():
        mapped = SYMPTOM_FINDINGS.get(label) or (
            ("urine_protein test",) if label in URINE_PROTEIN_LABELS else ()
        )
        print(f"  {n:4d}  {label!r} -> {', '.join(mapped) or '(no finding)'}")
    if not args.confirmed_by:
        print(
            "\nNot scoring: run with --confirmed-by once the clinical adviser has"
            " confirmed the mapping."
        )
        return 0
    res = evaluate(records)
    out = {"confirmed_by": args.confirmed_by, "records": len(records), "sha256": SHA256}
    for name, v in res.items():
        if hasattr(v, "k"):
            ci = v.ci()
            out[name] = {"k": v.k, "n": v.n, "point": v.point, "ci95": ci}
            print(
                f"{name}: {v.k}/{v.n}"
                + (f" = {v.point:.3f} [{ci[0]:.3f}, {ci[1]:.3f}]" if ci else "")
            )
        else:
            out[name] = v
            print(f"{name}: {v}")
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "uyo.report.json").write_text(json.dumps(out, indent=2) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
